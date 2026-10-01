#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera el archivo de datos PERSONAL (datos-juegos.json) a partir del Excel y, opcionalmente,
de una carpeta con fotos. Ese archivo NO se publica: se importa en la tablet desde el panel
oculto de la app (mantener pulsado el engranaje 2,5 s > "Importar datos").

Uso:
    python3 generar_datos.py JUEGOS.xlsx [carpeta_fotos] [salida.json]

Excel (pestañas):
    PREGUNTAS       B=pregunta, C=respuesta correcta, D..F=respuestas incorrectas (opcionales)
    SOPA DE LETRAS  fila 1 = nombre de cada tema; debajo, sus palabras
    FOTOS           NOMBRE | HERMANO (varias columnas) | PAREJA | HIJO (varias columnas) | FECHA
    RUTINA          (opcional) HORA | ACTIVIDAD | PREGUNTA | NOTA | ICONO

Fotos: NOMBRE_01.jpg, NOMBRE_02.jpg (actuales) y NOMBRE_03.jpg (de pequeño).
       El NOMBRE debe coincidir con el de la pestaña FOTOS (MIGUEL ANGEL = MIGUEL_ANGEL_01.jpg).
Requiere: pip install openpyxl pillow
"""
import sys, os, re, io, json, base64, unicodedata
import openpyxl
try:
    from PIL import Image, ImageOps
except ImportError:
    Image = None

DEFAULT_ROUTINE = [
    {"name": "Levantarse", "q": "¿A qué hora te levantas?", "note": "", "h": 7, "m": 30, "icon": "⏰"},
    {"name": "Desayuno", "q": "¿A qué hora es el desayuno?", "note": "", "h": 8, "m": 0, "icon": "☕"},
    {"name": "Centro de día", "q": "¿A qué hora bajas al centro de día?", "note": "si te toca", "h": 9, "m": 0, "icon": "\U0001F68C"},
    {"name": "Comer", "q": "¿A qué hora comes?", "note": "si estás en casa", "h": 13, "m": 0, "icon": "\U0001F374"},
    {"name": "Cenar", "q": "¿A qué hora cenas?", "note": "si estás en casa", "h": 20, "m": 30, "icon": "\U0001F319"},
]

def clean(s):
    s = str(s or '').upper().replace('Ñ', '\0')
    s = ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^A-Z\0]', '', s).replace('\0', 'Ñ')

def fmt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()

def sheet(wb, *names):
    for ws in wb.worksheets:
        if ws.title.strip().upper() in [n.upper() for n in names]:
            return ws
    return None

def read_questions(ws):
    out = []
    for r in range(2, ws.max_row + 1):
        q, a = fmt(ws.cell(r, 2).value), fmt(ws.cell(r, 3).value)
        if not q or not a:
            continue
        wrong = [fmt(ws.cell(r, c).value) for c in (4, 5, 6)]
        out.append({"q": q, "a": a, "wrong": [w for w in wrong if w]})
    return out

def read_themes(ws):
    out = []
    for c in range(1, ws.max_column + 1):
        name = fmt(ws.cell(1, c).value)
        if not name:
            continue
        words = [fmt(ws.cell(r, c).value) for r in range(2, ws.max_row + 1)]
        out.append({"name": name, "words": [w for w in words if w]})
    return out

def read_people(ws):
    heads = {}
    for c in range(1, ws.max_column + 1):
        h = fmt(ws.cell(1, c).value).upper()
        if h:
            heads.setdefault(h, []).append(c)
    def cols(*names):
        res = []
        for n in names:
            res += heads.get(n, [])
        return res
    people = []
    for r in range(2, ws.max_row + 1):
        name = fmt(ws.cell(r, (cols('NOMBRE') or [2])[0]).value)
        if not name:
            continue
        get = lambda names: [fmt(ws.cell(r, c).value) for c in cols(*names) if fmt(ws.cell(r, c).value)]
        pareja = get(['PAREJA'])
        fecha = get(['FECHA'])
        people.append({"name": name, "bday": fecha[0] if fecha else "",
                       "partner": pareja[0] if pareja else "",
                       "children": get(['HIJO', 'HIJOS']),
                       "siblings": get(['HERMANO', 'HERMANOS']),
                       "photos": [None, None, None]})
    return people

def read_routine(ws):
    out = []
    for r in range(2, ws.max_row + 1):
        hora, act = fmt(ws.cell(r, 1).value), fmt(ws.cell(r, 2).value)
        if not hora or not act:
            continue
        m = re.match(r'(\d{1,2})[:.](\d{2})', hora)
        if not m:
            continue
        out.append({"name": act, "q": fmt(ws.cell(r, 3).value) or "¿A qué hora es: %s?" % act.lower(),
                    "note": fmt(ws.cell(r, 4).value), "h": int(m.group(1)), "m": int(m.group(2)),
                    "icon": fmt(ws.cell(r, 5).value) or "⏰"})
    return out

def photo_data_url(path, maxside=700):
    im = Image.open(path)
    im = ImageOps.exif_transpose(im).convert('RGB')
    im.thumbnail((maxside, maxside))
    buf = io.BytesIO()
    im.save(buf, 'JPEG', quality=82)
    return 'data:image/jpeg;base64,' + base64.b64encode(buf.getvalue()).decode('ascii')

def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    xlsx = sys.argv[1]
    folder = sys.argv[2] if len(sys.argv) > 2 else None
    out = sys.argv[3] if len(sys.argv) > 3 else 'datos-juegos.json'
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws_q, ws_s, ws_f, ws_r = sheet(wb, 'PREGUNTAS'), sheet(wb, 'SOPA DE LETRAS', 'SOPA'), sheet(wb, 'FOTOS'), sheet(wb, 'RUTINA')
    data = {"version": 1,
            "people": read_people(ws_f) if ws_f else [],
            "questions": read_questions(ws_q) if ws_q else [],
            "themes": read_themes(ws_s) if ws_s else [],
            "routine": (read_routine(ws_r) if ws_r else []) or DEFAULT_ROUTINE}
    placed, unmatched = 0, []
    if folder and os.path.isdir(folder):
        if Image is None:
            sys.exit('Falta pillow: pip install pillow')
        index = {clean(p['name']): p for p in data['people']}
        for fn in sorted(os.listdir(folder)):
            m = re.match(r'^(.*?)[\s_\-]*0?([1-3])\.(jpe?g|png|webp)$', fn, re.I)
            p = index.get(clean(m.group(1))) if m else None
            if not p:
                unmatched.append(fn)
                continue
            p['photos'][int(m.group(2)) - 1] = photo_data_url(os.path.join(folder, fn))
            placed += 1
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False)
    print('Archivo:', out, '(%d KB)' % (os.path.getsize(out) // 1024))
    print('Personas: %d | preguntas: %d | temas de sopa: %d | actividades: %d | fotos colocadas: %d'
          % (len(data['people']), len(data['questions']), len(data['themes']), len(data['routine']), placed))
    for p in data['people']:
        miss = [str(i + 1).zfill(2) for i in range(3) if not p['photos'][i]]
        if miss:
            print('  Sin foto %s: %s' % (','.join(miss), p['name']))
    for fn in unmatched:
        print('  Foto sin coincidencia:', fn)

if __name__ == '__main__':
    main()
